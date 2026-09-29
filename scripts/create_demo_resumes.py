import _bootstrap
from src.utils.config import DATA
from docx import Document
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.colors import HexColor
from xml.sax.saxutils import escape

def create():
    styles=getSampleStyleSheet();styles['Title'].textColor=HexColor('#10233F');styles['Heading2'].textColor=HexColor('#2563EB')
    for path in sorted((DATA/'resumes').glob('candidate_*.txt')):
        lines=path.read_text().splitlines();doc=Document();story=[]
        for i,line in enumerate(lines):
            heading=line in ['Skills','Experience','Projects','Education','Certifications','Domains']
            if i==0:doc.add_heading(line,0)
            elif heading:doc.add_heading(line,2)
            else:doc.add_paragraph(line)
            story.append(Paragraph(escape(line),styles['Title' if i==0 else 'Heading2' if heading else 'BodyText']));story.append(Spacer(1,8))
        doc.save(path.with_suffix('.docx'));SimpleDocTemplate(str(path.with_suffix('.pdf')),rightMargin=50,leftMargin=50,topMargin=45,bottomMargin=45).build(story)
    print('Created two fictional PDF and DOCX resume pairs.')
if __name__=='__main__':create()
