from pathlib import Path
from io import BytesIO
from pypdf import PdfReader
from docx import Document

def extract_text(source, name=None):
    name = name or getattr(source, 'name', str(source))
    suffix = Path(name).suffix.lower()
    if isinstance(source, (str, Path)):
        data = Path(source).read_bytes()
    elif hasattr(source, 'getvalue'):
        data = source.getvalue()
    else:
        # Streamlit reruns must not consume an uploaded document permanently.
        position = source.tell()
        try:
            source.seek(0)
            data = source.read()
        finally:
            source.seek(position)
    if len(data) > 20 * 1024 * 1024: raise ValueError('File exceeds 20 MB limit.')
    if suffix == '.pdf':
        reader = PdfReader(BytesIO(data))
        if reader.is_encrypted: raise ValueError('Encrypted PDF is unsupported.')
        text = '\n'.join(p.extract_text() or '' for p in reader.pages)
    elif suffix == '.docx':
        doc = Document(BytesIO(data))
        text = '\n'.join([p.text for p in doc.paragraphs] + [' | '.join(c.text for c in row.cells) for t in doc.tables for row in t.rows])
    elif suffix == '.txt': text = data.decode('utf-8-sig')
    else: raise ValueError('Supported formats: PDF, DOCX and TXT (job descriptions).')
    if not text.strip(): raise ValueError('No extractable text. Scanned documents need OCR before upload.')
    return text.strip()
